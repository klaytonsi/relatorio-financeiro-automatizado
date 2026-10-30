"""Generate a formatted Excel report from the consolidated table.

Gera um relatório Excel formatado a partir da tabela consolidada.
"""

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# EN: Column names and allowed categories come from config.py (single source).
# PT: Nomes de colunas e categorias permitidas vêm do config.py (fonte única).
from config import (
    AMOUNT_COL,
    CATEGORIES,
    CATEGORY_COL,
    DATE_COL,
    ERROR_TYPE_COL,
)

# EN: Input and output paths, relative to the project root.
# PT: Caminhos de entrada e saída, relativos à raiz do projeto.
INPUT_FILE = Path("data/processed/consolidated.csv")
ERRORS_FILE = Path("reports/validation_errors.csv")
OUT_FILE = Path("reports/relatorio_financeiro.xlsx")

# EN: Headers of the report tables (what the reader sees in Excel).
# PT: Cabeçalhos das tabelas do relatório (o que o leitor vê no Excel).
MONTH_HEADER = "Mês"
CATEGORY_HEADER = "Categoria"
ENTRIES_HEADER = "Lançamentos"
TOTAL_HEADER = "Total (R$)"
ERROR_HEADER = "Tipo de erro"
QUANTITY_HEADER = "Quantidade"

# EN: Visual style. Arial is a safe, professional font in any Excel.
# PT: Estilo visual. Arial é uma fonte profissional e segura em qualquer Excel.
MONEY_FORMAT = '"R$" #,##0.00'
BASE_FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", start_color="1F4E78")


def load_valid_rows(path: Path = INPUT_FILE) -> tuple[pd.DataFrame, int]:
    """Load the consolidated table and keep only rows safe to sum.

    Carrega a tabela consolidada e mantém só as linhas seguras para somar.
    """
    df = pd.read_csv(path)
    missing = [c for c in (DATE_COL, CATEGORY_COL, AMOUNT_COL) if c not in df]
    if missing:
        raise ValueError(f"Missing columns / Colunas ausentes: {missing}")

    # EN: Convert types. Bad values become NaT/NaN instead of raising errors.
    # PT: Converte os tipos. Valores ruins viram NaT/NaN em vez de dar erro.
    df[DATE_COL] = pd.to_datetime(df[DATE_COL], errors="coerce")
    df[AMOUNT_COL] = pd.to_numeric(df[AMOUNT_COL], errors="coerce")
    df[CATEGORY_COL] = df[CATEGORY_COL].astype("string").str.strip()

    # EN: A row is valid if it has a date, an allowed category and amount >= 0.
    #     ge(0) is False for NaN and isin() is False for empty categories,
    #     so empty values are excluded too.
    # PT: Linha válida tem data, categoria permitida e valor >= 0.
    #     ge(0) é False para NaN e isin() é False para categoria vazia,
    #     então valores vazios também ficam de fora.
    is_valid = (
        df[DATE_COL].notna()
        & df[CATEGORY_COL].isin(CATEGORIES)
        & df[AMOUNT_COL].ge(0)
    )
    return df[is_valid].copy(), int((~is_valid).sum())


def load_errors(path: Path = ERRORS_FILE) -> pd.DataFrame:
    """Load the validation error report.

    Carrega o relatório de erros da validação.
    """
    return pd.read_csv(path)


def summarize_by_month(df: pd.DataFrame) -> pd.DataFrame:
    """Count entries and sum amounts per month (YYYY-MM).

    Conta lançamentos e soma valores por mês (AAAA-MM).
    """
    # EN: "YYYY-MM" text sorts in the correct chronological order.
    # PT: Texto "AAAA-MM" ordena na ordem cronológica correta.
    month = df[DATE_COL].dt.to_period("M").astype(str).rename(MONTH_HEADER)
    grouped = df.groupby(month)[AMOUNT_COL].agg(["count", "sum"])
    grouped.columns = [ENTRIES_HEADER, TOTAL_HEADER]
    return grouped.reset_index()


def summarize_by_category(df: pd.DataFrame) -> pd.DataFrame:
    """Count entries and sum amounts per category, largest total first.

    Conta lançamentos e soma valores por categoria, maior total primeiro.
    """
    category = df[CATEGORY_COL].rename(CATEGORY_HEADER)
    grouped = df.groupby(category)[AMOUNT_COL].agg(["count", "sum"])
    grouped.columns = [ENTRIES_HEADER, TOTAL_HEADER]
    return grouped.sort_values(TOTAL_HEADER, ascending=False).reset_index()


def summarize_errors(errors: pd.DataFrame) -> pd.DataFrame:
    """Count validation errors per type.

    Conta os erros de validação por tipo.
    """
    counts = errors[ERROR_TYPE_COL].value_counts()
    return counts.rename_axis(ERROR_HEADER).reset_index(name=QUANTITY_HEADER)


def style_sheet(
    ws: Worksheet,
    money_headers: tuple[str, ...] = (),
    total_headers: tuple[str, ...] = (),
) -> None:
    """Format one sheet: header, widths, money format and a total row.

    Formata uma aba: cabeçalho, larguras, formato de moeda e linha de total.
    """
    headers = [cell.value for cell in ws[1]]
    last_data_row = ws.max_row

    # EN: Column width from the longest text; money needs extra room for "R$".
    # PT: Largura pelo maior texto; moeda precisa de espaço extra para "R$".
    for idx, name in enumerate(headers, start=1):
        longest = max(len(str(c.value or "")) for c in ws[get_column_letter(idx)])
        extra = 8 if name in money_headers else 2
        ws.column_dimensions[get_column_letter(idx)].width = longest + extra

    # EN: Total row uses SUM formulas, so the sheet recalculates if edited.
    # PT: A linha de total usa fórmulas SUM, então recalcula se for editada.
    if total_headers and last_data_row >= 2:
        total_row = last_data_row + 1
        ws.cell(total_row, 1, "Total")
        for name in total_headers:
            col = headers.index(name) + 1
            letter = get_column_letter(col)
            ws.cell(total_row, col, f"=SUM({letter}2:{letter}{last_data_row})")

    # EN: Apply fonts and number formats to every cell.
    # PT: Aplica fontes e formatos numéricos em todas as células.
    for row in ws.iter_rows(min_row=2):
        is_total = row[0].value == "Total"
        for cell in row:
            cell.font = BOLD_FONT if is_total else BASE_FONT
            if headers[cell.column - 1] in money_headers:
                cell.number_format = MONEY_FORMAT

    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"


def build_report(
    input_file: Path = INPUT_FILE,
    errors_file: Path = ERRORS_FILE,
    out_file: Path = OUT_FILE,
) -> tuple[Path, int]:
    """Build the Excel report and return its path and excluded row count.

    Monta o relatório Excel e devolve o caminho e o nº de linhas excluídas.
    """
    valid, excluded = load_valid_rows(input_file)
    errors = load_errors(errors_file)

    # EN: (sheet name, table, money columns, columns that get a total)
    # PT: (nome da aba, tabela, colunas de moeda, colunas que recebem total)
    sheets = [
        (
            "Por mês",
            summarize_by_month(valid),
            (TOTAL_HEADER,),
            (ENTRIES_HEADER, TOTAL_HEADER),
        ),
        (
            "Por categoria",
            summarize_by_category(valid),
            (TOTAL_HEADER,),
            (ENTRIES_HEADER, TOTAL_HEADER),
        ),
        ("Erros", summarize_errors(errors), (), (QUANTITY_HEADER,)),
    ]

    out_file.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(out_file, engine="openpyxl") as writer:
        for name, table, money, totals in sheets:
            table.to_excel(writer, sheet_name=name, index=False)
            style_sheet(writer.sheets[name], money, totals)
    return out_file, excluded


def main() -> None:
    """Check inputs, build the report and print a short summary.

    Confere as entradas, monta o relatório e imprime um resumo curto.
    """
    # EN: Fail with a clear message if a previous step was not run.
    # PT: Falha com mensagem clara se uma etapa anterior não foi rodada.
    for required in (INPUT_FILE, ERRORS_FILE):
        if not required.exists():
            raise FileNotFoundError(
                f"{required} not found. Run consolidate.py and validate.py "
                f"first. / {required} não encontrado. Rode o consolidate.py "
                f"e o validate.py primeiro."
            )

    path, excluded = build_report()
    print(f"Relatório salvo / Report saved: {path}")
    print(f"{excluded} linhas fora dos resumos / rows left out of summaries")


if __name__ == "__main__":
    main()
      
